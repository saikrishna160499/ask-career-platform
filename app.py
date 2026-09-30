import os
import io
import re
import requests
import streamlit as st
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader
from bs4 import BeautifulSoup
from concurrent.futures import ThreadPoolExecutor, as_completed
from google import genai
from google.genai import types

# ---------------------------------------------------------
# Page Configuration & Styling
# ---------------------------------------------------------
st.set_page_config(
    page_title="Ask Career Portal",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
    <style>
    .main { background-color: #0e1117; }
    .job-card {
        background-color: #1e222b;
        border: 1px solid #30363d;
        border-radius: 8px;
        padding: 16px;
        margin-bottom: 14px;
    }
    .platform-badge {
        display: inline-block;
        padding: 2px 10px;
        font-size: 12px;
        font-weight: 600;
        border-radius: 12px;
        background-color: #238636;
        color: #ffffff;
        margin-right: 6px;
    }
    .emp-badge {
        display: inline-block;
        padding: 2px 10px;
        font-size: 12px;
        font-weight: 600;
        border-radius: 12px;
        background-color: #1f6feb;
        color: #ffffff;
    }
    </style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# Dynamic Secrets & Environment Resolver
# ---------------------------------------------------------
def get_secret(key_name: str, default: str = "") -> str:
    """Retrieves secret keys from Streamlit secrets or OS environment variables."""
    if hasattr(st, "secrets") and key_name in st.secrets:
        return st.secrets[key_name]
    return os.environ.get(key_name, default)

GEMINI_API_KEY = get_secret("GEMINI_API_KEY")
SERPAPI_KEY = get_secret("SERPAPI_KEY")
RAPIDAPI_KEY = get_secret("RAPIDAPI_KEY")

TARGET_BOARDS = [
    "LinkedIn",
    "Dice",
    "Indeed",
    "Monster",
    "Glassdoor",
    "ZipRecruiter",
    "Robert Half",
    "CareerBuilder",
    "SimplyHired"
]

EMPLOYMENT_TYPES = ["All Types", "Full-Time", "Contract / C2C", "Part-Time", "Internship"]

# ---------------------------------------------------------
# Multi-Board Ingestion Engine
# ---------------------------------------------------------
def fetch_board_specific_jobs(role: str, location: str, board: str, job_type: str) -> list:
    """Queries Google Jobs API via SerpApi with targeted search syntax per board & job type."""
    if not SERPAPI_KEY:
        return []
    
    # Construct search operators tailored per platform
    domain_map = {
        "LinkedIn": "linkedin.com/jobs",
        "Dice": "dice.com",
        "Indeed": "indeed.com",
        "Monster": "monster.com",
        "Glassdoor": "glassdoor.com",
        "ZipRecruiter": "ziprecruiter.com",
        "Robert Half": "roberthalf.com",
        "CareerBuilder": "careerbuilder.com",
        "SimplyHired": "simplyhired.com"
    }
    
    board_domain = domain_map.get(board, f"{board.lower().replace(' ', '')}.com")
    
    type_query = ""
    if job_type == "Contract / C2C":
        type_query = "contract OR c2c OR \"corp-to-corp\" OR 1099 OR temp"
    elif job_type == "Full-Time":
        type_query = "\"full time\" OR \"full-time\""
    elif job_type == "Part-Time":
        type_query = "\"part time\" OR \"part-time\""
    elif job_type == "Internship":
        type_query = "internship OR intern"

    query_str = f"{role} {location} {type_query} site:{board_domain}".strip()

    params = {
        "engine": "google_jobs",
        "q": query_str,
        "location": location,
        "hl": "en",
        "api_key": SERPAPI_KEY,
        "num": 10
    }
    
    board_results = []
    try:
        response = requests.get("https://serpapi.com/search", params=params, timeout=7)
        if response.status_code == 200:
            data = response.json()
            for item in data.get("jobs_results", []):
                # Detect direct application link
                apply_link = ""
                for opt in item.get("apply_options", []):
                    if opt.get("link"):
                        apply_link = opt.get("link")
                        break
                
                # Detect or infer employment label
                detected_type = "Full-Time"
                desc_lower = (item.get("description", "") + " " + item.get("title", "")).lower()
                if any(k in desc_lower for k in ["contract", "c2c", "corp to corp", "1099", "temp"]):
                    detected_type = "Contract"
                elif "part time" in desc_lower or "part-time" in desc_lower:
                    detected_type = "Part-Time"
                
                board_results.append({
                    "title": item.get("title", role),
                    "company": item.get("company_name", "Confidential Hiring Client"),
                    "location": item.get("location", location),
                    "platform": board,
                    "employment_type": detected_type,
                    "link": apply_link or item.get("share_link", "https://google.com"),
                    "description": item.get("description", "No detailed summary provided.")
                })
    except Exception:
        pass
    
    return board_results

def fetch_all_multiboard_jobs(role: str, location: str, job_type: str) -> list:
    """Executes parallel threads across all 9 target boards."""
    aggregated = []
    with ThreadPoolExecutor(max_workers=9) as executor:
        futures = {
            executor.submit(fetch_board_specific_jobs, role, location, b, job_type): b 
            for b in TARGET_BOARDS
        }
        for future in as_completed(futures):
            try:
                res = future.result()
                if res:
                    aggregated.extend(res)
            except Exception:
                continue

    # Deduplicate entries by unique (Title + Company) key
    seen = set()
    deduped = []
    for j in aggregated:
        unique_key = f"{j['title'].lower().strip()}_{j['company'].lower().strip()}"
        if unique_key not in seen:
            seen.add(unique_key)
            deduped.append(j)
            
    return deduped

# ---------------------------------------------------------
# Document Extraction & Generation Utilities
# ---------------------------------------------------------
def extract_text_from_file(uploaded_file) -> str:
    """Extracts raw text content from PDF or DOCX file objects."""
    if uploaded_file.name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    elif uploaded_file.name.endswith(".docx"):
        doc = Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs if p.text])
    return ""

def generate_formatted_docx(resume_text: str) -> io.BytesIO:
    """Builds a clean, professional ATS-standard DOCX document."""
    doc = Document()
    
    # 0.5-inch compact professional margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    lines = resume_text.strip().split("\n")
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
            
        # Section Header Detection
        if line.startswith("# ") or (line.isupper() and len(line) < 35):
            h_text = line.replace("# ", "").strip()
            p = doc.add_paragraph()
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("* ") or line.startswith("- ") or line.startswith("• "):
            bullet_text = line.lstrip("*-• ").strip()
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(bullet_text)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# Sidebar Navigation
# ---------------------------------------------------------
with st.sidebar:
    st.title("⚡ Ask Career Portal")
    st.caption("AI-Powered Job Aggregation & ATS Resume Suite")
    st.divider()
    active_tab = st.radio(
        "Navigation",
        ["💼 Live Multi-Board Search", "📄 ATS Resume Re-Architect", "✉️ Cold Outreach Pitch"],
        index=0
    )
    st.divider()
    st.markdown("### Integrated Data Engines")
    for b in TARGET_BOARDS:
        st.markdown(f"- ✅ **{b}**")

# ---------------------------------------------------------
# View 1: Live Multi-Board Search
# ---------------------------------------------------------
if active_tab == "💼 Live Multi-Board Search":
    st.header("💼 Live Multi-Board Market Search")
    st.write("Concurrently scrape and query full-time, contract, and remote listings across all primary job portals.")

    c1, c2, c3 = st.columns([3, 2, 2])
    with c1:
        target_role = st.text_input("Target Job Title / Domain", value="Data Engineer")
    with c2:
        target_location = st.text_input("Target Location", value="United States")
    with c3:
        target_type = st.selectbox("Employment Contract Model", EMPLOYMENT_TYPES, index=0)

    if st.button("🔎 Run Aggregated Multi-Board Query", use_container_width=True):
        if not SERPAPI_KEY:
            st.error("Missing `SERPAPI_KEY`. Please configure your key in Render Environment Variables or `.streamlit/secrets.toml`.")
        else:
            with st.spinner("Connecting to Dice, LinkedIn, Indeed, Monster, Glassdoor, ZipRecruiter, Robert Half..."):
                found_jobs = fetch_all_multiboard_jobs(target_role, target_location, target_type)
                st.session_state["portal_jobs"] = found_jobs
                if not found_jobs:
                    st.warning("No listings returned. Verify your SerpApi quota or expand search parameters.")

    if "portal_jobs" in st.session_state and st.session_state["portal_jobs"]:
        jobs = st.session_state["portal_jobs"]
        st.success(f"Successfully aggregated **{len(jobs)} unique positions**.")

        # Metric Distribution Bar
        st.subheader("Platform Distribution Breakdown")
        dist_cols = st.columns(len(TARGET_BOARDS))
        counts = {b: 0 for b in TARGET_BOARDS}
        for j in jobs:
            counts[j["platform"]] = counts.get(j["platform"], 0) + 1
            
        for i, board in enumerate(TARGET_BOARDS):
            with dist_cols[i]:
                st.metric(label=board, value=counts[board])

        st.divider()

        # Dynamic Filtering Bar
        fcol1, fcol2 = st.columns([2, 2])
        with fcol1:
            selected_board = st.selectbox("Filter Display by Portal", ["All Portals"] + TARGET_BOARDS)
        with fcol2:
            selected_emp = st.selectbox("Filter Display by Type", ["All Types", "Full-Time", "Contract", "Part-Time"])

        filtered = jobs
        if selected_board != "All Portals":
            filtered = [j for j in filtered if j["platform"] == selected_board]
        if selected_emp != "All Types":
            filtered = [j for j in filtered if j["employment_type"] == selected_emp]

        st.write(f"Showing **{len(filtered)}** filtered results:")

        for job in filtered:
            with st.container():
                st.markdown(f"""
                <div class="job-card">
                    <span class="platform-badge">{job['platform']}</span>
                    <span class="emp-badge">{job['employment_type']}</span>
                    <h3 style="margin-top: 8px; margin-bottom: 4px;">{job['title']}</h3>
                    <p style="color: #8b949e; margin-bottom: 8px;"><strong>{job['company']}</strong> — 📍 {job['location']}</p>
                    <p style="font-size: 14px; line-height: 1.5;">{job['description'][:300]}...</p>
                </div>
                """, unsafe_allow_html=True)
                
                c_btn1, c_btn2 = st.columns([2, 5])
                with c_btn1:
                    st.link_button(f"🚀 Apply on {job['platform']}", job["link"], use_container_width=True)
                with c_btn2:
                    if st.button("📋 Stage for Resume Tailoring", key=f"stage_{job['title']}_{job['company']}"):
                        st.session_state["staged_jd"] = f"Title: {job['title']}\nCompany: {job['company']}\n\n{job['description']}"
                        st.info("Job description staged into ATS Resume Builder tab.")
                st.write("")

# ---------------------------------------------------------
# View 2: ATS Resume Re-Architect
# ---------------------------------------------------------
elif active_tab == "📄 ATS Resume Re-Architect":
    st.header("📄 ATS Resume Re-Architect & Builder")
    st.write("Align candidate project histories with target job descriptions using Gemini 2.5.")

    uploaded_doc = st.file_uploader("Upload Current Master Resume (.docx or .pdf)", type=["docx", "pdf"])
    
    staged_default = st.session_state.get("staged_jd", "")
    target_jd = st.text_area("Target Job Description", value=staged_default, height=220)

    if st.button("🚀 Re-Architect Resume & Generate DOCX", use_container_width=True):
        if not GEMINI_API_KEY:
            st.error("Missing `GEMINI_API_KEY`. Please configure your key in Render Environment Variables or `.streamlit/secrets.toml`.")
        elif not uploaded_doc or not target_jd:
            st.error("Please provide both an uploaded resume file and a target job description.")
        else:
            with st.spinner("Analyzing skill taxonomies and generating tailored experience points..."):
                try:
                    resume_text = extract_text_from_file(uploaded_doc)
                    client = genai.Client(api_key=GEMINI_API_KEY)
                    
                    prompt = f"""
You are an expert Executive Resume Writer and ATS Optimization Specialist.
Re-architect the candidate's resume to match the target Job Description precisely.

STRICT INSTRUCTIONS:
1. Preserve candidate contact details, certifications, and educational degrees.
2. Update the Professional Summary (concise, high-impact, keyword-rich).
3. Align the Technical Skills Inventory to explicitly feature the required cloud platforms, frameworks, tools, and languages found in the JD.
4. For the candidate's primary/recent projects, generate exactly 15 detailed, high-impact bullet points each. 
   - Every bullet must follow the Action Verb + Context/Challenge + Technical Architecture + Measurable Business Outcome formula.
   - Bullets must be dense (1.5 to 2 lines in length).
   - Append an explicit Environment line at the conclusion of each role (e.g., Environment: Tool1, Tool2, Cloud, Framework).
5. Output format must be clean markdown with headers marked as # PROFESSIONAL SUMMARY, # TECHNICAL SKILLS, # PROFESSIONAL EXPERIENCE, etc.

Target Job Description:
{target_jd}

Candidate Current Resume:
{resume_text}
"""
                    response = client.models.generate_content(
                        model="gemini-2.5-flash",
                        contents=prompt
                    )
                    
                    st.session_state["tailored_resume_text"] = response.text
                    st.success("Resume re-architected successfully.")
                except Exception as e:
                    st.error(f"Generation failure: {str(e)}")

    if "tailored_resume_text" in st.session_state:
        st.divider()
        st.subheader("Tailored Resume Preview")
        st.markdown(st.session_state["tailored_resume_text"])
        
        docx_file = generate_formatted_docx(st.session_state["tailored_resume_text"])
        st.download_button(
            label="📥 Download Tailored Resume (.docx)",
            data=docx_file,
            file_name="Tailored_Professional_Resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )

# ---------------------------------------------------------
# View 3: Cold Outreach Pitch
# ---------------------------------------------------------
elif active_tab == "✉️ Cold Outreach Pitch":
    st.header("✉️ Cold Outreach & Recruiter Pitch Generator")
    st.write("Generate tailored hiring manager emails and LinkedIn connection notes.")

    col_a, col_b = st.columns(2)
    with col_a:
        lead_role = st.text_input("Role Title", value="Senior Data Engineer")
        hiring_firm = st.text_input("Company Name", value="Target Organization")
    with col_b:
        recruiter_name = st.text_input("Recipient / Recruiter Name", value="Hiring Team")
        contract_type = st.selectbox("Engagement Type", ["Full-Time W2", "Contract / C2C (Corp-to-Corp)", "General Ingestion"])

    core_strengths = st.text_area("Key Strengths & Differentiators to Highlight", 
                                  value="Cloud database migration, distributed data pipelines, enterprise analytics, automated QA frameworks.")

    if st.button("⚡ Generate Outreach Messages", use_container_width=True):
        if not GEMINI_API_KEY:
            st.error("Missing `GEMINI_API_KEY`.")
        else:
            with st.spinner("Drafting executive pitch templates..."):
                try:
                    client = genai.Client(api_key=GEMINI_API_KEY)
                    outreach_prompt = f"""
Write two targeted professional outreach messages for a candidate reaching out for a {lead_role} role ({contract_type}) at {hiring_firm}.
Addressed to: {recruiter_name}
Core Technical Highlights: {core_strengths}

Generate:
1. A 300-character max LinkedIn Connection Request Note.
2. A 3-paragraph high-conversion direct Recruiter Email with a compelling Subject Line, highlighting quantifiable delivery and zero-fluff expertise.
"""
                    outreach_res = client.models.generate_content(
                        model="gemini-2.5-flash",
                        contents=outreach_prompt
                    )
                    st.markdown(outreach_res.text)
                except Exception as e:
                    st.error(f"Failed to generate outreach pitch: {str(e)}")